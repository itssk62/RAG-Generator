"""Download the two public-domain eval corpora and normalise them into the supported formats.

tech/        NIST Cybersecurity Framework 2.0 (PDF, US Gov public domain)
             PEP 8 (reStructuredText, public domain) - ingested through the `.rst -> TextParser` config mapping
             PEP 257 (converted to Markdown, public domain)
humanities/  Sun Tzu, The Art of War, tr. Lionel Giles (TXT, Project Gutenberg #132, public domain)
             Epictetus, The Enchiridion, tr. T. W. Higginson (built as DOCX, Gutenberg #45109, public domain)

Usage: python eval/fetch_corpora.py
"""

from __future__ import annotations

import re
import sys
import urllib.request
from pathlib import Path

import docx

ROOT = Path(__file__).parent / "corpora"
SOURCES = {
    "nist_csf": "https://nvlpubs.nist.gov/nistpubs/CSWP/NIST.CSWP.29.pdf",
    "pep8": "https://raw.githubusercontent.com/python/peps/main/peps/pep-0008.rst",
    "pep257": "https://raw.githubusercontent.com/python/peps/main/peps/pep-0257.rst",
    # Project Gutenberg texts via the GITenberg mirror on GitHub
    "art_of_war": "https://raw.githubusercontent.com/GITenberg/The-Art-of-War_132/master/132.txt",
    "enchiridion": "https://raw.githubusercontent.com/GITenberg/The-Enchiridion_45109/master/45109-0.txt",
}


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "rag-generator-eval/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def strip_gutenberg(text: str) -> str:
    text = text.replace("\r\n", "\n")
    start = re.search(r"\*\*\* ?START OF (THE|THIS) PROJECT GUTENBERG.*?\*\*\*", text)
    end = re.search(r"\*\*\* ?END OF (THE|THIS) PROJECT GUTENBERG", text)
    if start:
        text = text[start.end() : end.start() if end else None]
    elif re.search(r"End of (the )?Project Gutenberg", text):
        text = text[: re.search(r"End of (the )?Project Gutenberg", text).start()]
    return text.strip() + "\n"


def rst_to_markdown(text: str) -> str:
    """Just enough RST -> Markdown for PEP 257: underlined titles become `#` headings."""
    lines = text.replace("\r\n", "\n").split("\n")
    out, levels, i = [], [], 0
    while i < len(lines):
        line = lines[i]
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if line.strip() and re.fullmatch(r"([=\-~^])\1{2,}", nxt.strip() or "x") and len(nxt.strip()) >= len(line.strip()) - 1:
            char = nxt.strip()[0]
            if char not in levels:
                levels.append(char)
            out.append("#" * (levels.index(char) + 1) + " " + line.strip())
            i += 2
            continue
        out.append(re.sub(r"``([^`]+)``", r"`\1`", line))
        i += 1
    return "\n".join(out)


def enchiridion_docx(text: str, path: Path) -> None:
    """Numbered sections ("I.", "II.", ...) become Heading 1 paragraphs."""
    body = strip_gutenberg(text)
    d = docx.Document()
    d.add_heading("The Enchiridion", level=0)
    for para in re.split(r"\n\s*\n", body):
        para = " ".join(para.split())
        if not para:
            continue
        if re.fullmatch(r"[IVXLC]+\.?", para):
            d.add_heading(f"Section {para.rstrip('.')}", level=1)
        else:
            d.add_paragraph(para)
    d.save(path)


def main() -> int:
    tech, hum = ROOT / "tech", ROOT / "humanities"
    tech.mkdir(parents=True, exist_ok=True)
    hum.mkdir(parents=True, exist_ok=True)
    (tech / "nist_csf_2.0.pdf").write_bytes(fetch(SOURCES["nist_csf"]))
    (tech / "pep8_style_guide.rst").write_bytes(fetch(SOURCES["pep8"]))
    (tech / "pep257_docstrings.md").write_text(rst_to_markdown(fetch(SOURCES["pep257"]).decode()), encoding="utf-8")
    (hum / "art_of_war.txt").write_text(strip_gutenberg(fetch(SOURCES["art_of_war"]).decode("utf-8", "replace")), encoding="utf-8")
    enchiridion_docx(fetch(SOURCES["enchiridion"]).decode("utf-8", "replace"), hum / "enchiridion.docx")
    for f in sorted(ROOT.rglob("*.*")):
        print(f"{f.relative_to(ROOT)}  {f.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
