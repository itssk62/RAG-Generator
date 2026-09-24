from __future__ import annotations

import docx
import pymupdf
import pytest

from app.core.interfaces import Block
from app.ingestion.parsers.docx import DocxParser
from app.ingestion.parsers.markdown import MarkdownParser
from app.ingestion.parsers.pdf import PdfParser
from app.ingestion.parsers.registry import ParserRegistry, UnsupportedFormat
from app.ingestion.parsers.text import TextParser


def test_text_parser_detects_underline_caps_and_numbered_headings():
    text = """Style Guide
===========

Intro paragraph that is
wrapped over two lines.

Code Layout
-----------

Use four spaces.

III. ATTACK BY STRATAGEM

In the practical art of war, the best thing of all is to take the enemy's country whole.

This line ends with a period.
"""
    blocks = TextParser().parse_text(text)
    assert blocks == [
        Block("Style Guide", heading_level=1),
        Block("Intro paragraph that is wrapped over two lines."),
        Block("Code Layout", heading_level=2),
        Block("Use four spaces."),
        Block("III. ATTACK BY STRATAGEM", heading_level=1),
        Block("In the practical art of war, the best thing of all is to take the enemy's country whole."),
        Block("This line ends with a period."),
    ]


def test_text_parser_overline_titles_and_numbered_sentences():
    text = "------------\nINTRODUCTION\n\n27.  When there is much running about\n\n2.1 Scope Of Work\n\nBody."
    assert [(b.text, b.heading_level) for b in TextParser().parse_text(text)] == [
        ("INTRODUCTION", 1),
        ("27. When there is much running about", 0),
        ("2.1 Scope Of Work", 1),
        ("Body.", 0),
    ]


def test_markdown_parser_headings_front_matter_and_fences():
    text = """---
title: x
---
# Guide

Some text
continues here.

## Install

```bash
pip install x

echo done
```
"""
    blocks = MarkdownParser().parse_text(text)
    assert blocks[0] == Block("Guide", heading_level=1)
    assert blocks[1] == Block("Some text\ncontinues here.")
    assert blocks[2] == Block("Install", heading_level=2)
    assert blocks[3].text.startswith("```bash") and "echo done" in blocks[3].text  # blank line inside fence kept


def test_docx_parser_uses_heading_styles(tmp_path):
    d = docx.Document()
    d.add_heading("Enchiridion", level=0)
    d.add_heading("Chapter 1", level=1)
    d.add_paragraph("Some things are in our control and others not.")
    d.add_paragraph("")
    d.add_heading("Section 1.1", level=2)
    d.add_paragraph("Things in our control are opinion and pursuit.")
    path = tmp_path / "e.docx"
    d.save(path)

    assert DocxParser().parse(path) == [
        Block("Enchiridion", heading_level=1),
        Block("Chapter 1", heading_level=1),
        Block("Some things are in our control and others not."),
        Block("Section 1.1", heading_level=2),
        Block("Things in our control are opinion and pursuit."),
    ]


def test_pdf_parser_keeps_pages_and_finds_headings_by_font_size(tmp_path):
    pdf = pymupdf.open()
    body = "The framework organizes outcomes into functions that apply to any organization."
    for page_no, title in enumerate(["Govern", "Identify"], start=1):
        page = pdf.new_page()
        page.insert_text((72, 72), title, fontsize=20)
        for i in range(4):
            page.insert_text((72, 120 + 20 * i), f"{body} Line {i}.", fontsize=10)
        page.insert_text((300, 800), str(page_no), fontsize=10)  # page-number footer
    path = tmp_path / "doc.pdf"
    pdf.save(path)

    blocks = PdfParser().parse(path)
    headings = [(b.text, b.page) for b in blocks if b.heading_level]
    assert headings == [("Govern", 1), ("Identify", 2)]
    assert {b.page for b in blocks if not b.heading_level} == {1, 2}
    assert not any(b.text in {"1", "2"} for b in blocks)


def test_pdf_parser_bold_headings_and_running_headers(tmp_path):
    pdf = pymupdf.open()
    for page_no in range(1, 4):
        page = pdf.new_page()
        page.insert_text((72, 40), "ACME REPORT 2024", fontsize=9)  # running header
        page.insert_text((72, 80), f"{page_no}. Section Title", fontsize=11, fontname="helvetica-bold")
        page.insert_text((72, 110), f"Fig. {page_no}. A caption", fontsize=11, fontname="helvetica-bold")
        for i in range(5):
            page.insert_text((72, 140 + 18 * i), f"Body text line {i} of page {page_no} for the report.", fontsize=11)
    path = tmp_path / "bold.pdf"
    pdf.save(path)

    blocks = PdfParser().parse(path)
    assert [(b.text, b.page) for b in blocks if b.heading_level] == [("1. Section Title", 1), ("2. Section Title", 2), ("3. Section Title", 3)]
    assert not any("ACME REPORT" in b.text for b in blocks)
    assert any(b.text.startswith("Fig. 1") and b.heading_level == 0 for b in blocks)


def test_registry_maps_extensions_from_config(tmp_path):
    reg = ParserRegistry(
        {
            ".TXT": "app.ingestion.parsers.text:TextParser",
            ".csv": {"class": "app.ingestion.parsers.text:TextParser", "options": {"encoding": "latin-1"}},
        }
    )
    assert reg.extensions == [".csv", ".txt"]
    assert reg.supports("notes.txt") and reg.supports("DATA.CSV") and not reg.supports("a.pdf")
    assert reg.for_file("data.csv").encoding == "latin-1"
    assert reg.for_file("a.txt") is reg.for_file("b.txt")
    with pytest.raises(UnsupportedFormat, match="supported: .csv, .txt"):
        reg.for_file("slides.pptx")
